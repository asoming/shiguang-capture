import io
import zipfile
import pytest
from shiguang_capture.structured import (
    TableData, code_from_blocks, table_from_blocks, table_risks,
    table_clipboard, markdown_table, xlsx_bytes, code_block,
)


def block(text,x,y,width=None):
    width = width or len(text)*10
    return {'text':text,'box':[[x,y],[x+width,y],[x+width,y+20],[x,y+20]],'confidence':.9}


def test_visible_code_indentation_and_symbols_preserved():
    blocks=[block('if ok:',20,10),block('print("00123")',60,40),block('else:',20,70),block('pass',60,100)]
    assert code_from_blocks(blocks) == 'if ok:\n    print("00123")\nelse:\n    pass'


def test_fragmented_code_on_same_line():
    assert code_from_blocks([block('x =',10,10),block('value',50,10)]) == 'x = value'


def test_empty_cells_keep_grid_positions():
    boxes=[[(0,0,100,50),(100,0,200,50)],[(0,50,100,100),(100,50,200,100)]]
    table=table_from_blocks(boxes,[block('00123',5,10),block('last',105,65)])
    assert table.cells == [['00123',''],['','last']]


@pytest.mark.parametrize('value',['=1+1',' +2','-12','@SUM(A1)','\ttext','multi\nline','\r=1'])
def test_risky_spreadsheet_clipboard_is_blocked(value):
    assert table_risks([['safe',value]]) == [(0,1)]
    with pytest.raises(ValueError):
        table_clipboard([['safe',value]])


def test_html_and_markdown_are_escaped():
    plain, rich=table_clipboard([['<img src="https://invalid.test/x">','A&B']])
    assert '<img' not in rich and '&lt;img' in rich
    assert 'A&amp;B' in rich
    md=markdown_table([['a|b','<script>'],['x\ny','&']])
    assert r'a\|b' in md and '&lt;script&gt;' in md and 'x<br>y' in md


def test_xlsx_strings_roundtrip_without_formulas():
    from openpyxl import load_workbook
    values=[['00123','12345678901234567890','2026-09-20','=HYPERLINK("https://invalid.test")'],['','-12','multi\nline','@SUM(A1)']]
    binary=xlsx_bytes(values)
    book=load_workbook(io.BytesIO(binary),data_only=False)
    for r,row in enumerate(values,1):
        for c,value in enumerate(row,1):
            cell=book.active.cell(r,c)
            assert (cell.value or '') == value
            assert cell.data_type != 'f'
            assert cell.number_format == '@'
    with zipfile.ZipFile(io.BytesIO(binary)) as archive:
        assert b'<f>' not in archive.read('xl/worksheets/sheet1.xml')
        assert not any('externalLinks' in name for name in archive.namelist())


def test_code_fences_cannot_be_closed_by_content():
    assert code_block('```').startswith('````\n')


def test_grid_detects_blank_cells_and_rejects_merged():
    cv2=pytest.importorskip('cv2')
    import numpy as np
    from shiguang_capture.structured import detect_grid
    image=np.full((180,330,3),255,np.uint8)
    for x in [10,110,210,310]: cv2.line(image,(x,10),(x,160),(0,0,0),2)
    for y in [10,60,110,160]: cv2.line(image,(10,y),(310,y),(0,0,0),2)
    data=cv2.imencode('.png',image)[1].tobytes()
    boxes,_=detect_grid(data)
    assert len(boxes)==3 and len(boxes[0])==3
    cv2.line(image,(110,13),(110,57),(255,255,255),4)
    with pytest.raises(ValueError):
        detect_grid(cv2.imencode('.png',image)[1].tobytes())


def test_xlsx_refuses_silent_truncation_and_control_characters():
    from shiguang_capture.structured import xlsx_bytes
    for value in ('a'*32768, 'bad\x00cell'):
        with pytest.raises(ValueError):
            xlsx_bytes([[value]])


def test_grid_cleaning_preserves_letter_stems():
    cv2 = pytest.importorskip('cv2')
    import numpy as np
    from shiguang_capture.structured import detect_grid
    img = np.full((340, 540, 3), 255, np.uint8)
    for x in (20, 270, 520): cv2.line(img, (x,20), (x,320), (0,0,0), 2)
    for y in (20,120,220,320): cv2.line(img, (20,y), (520,y), (0,0,0), 2)
    # Long, vertically aligned letter stems must not be counted as merged dividers.
    for y in (50,150,250): cv2.rectangle(img, (290,y), (294,y+30), (0,0,0), -1)
    ok, png = cv2.imencode('.png', img)
    boxes, clean = detect_grid(png.tobytes())
    restored = cv2.imdecode(np.frombuffer(clean,np.uint8),cv2.IMREAD_COLOR)
    assert len(boxes)==3 and len(boxes[0])==2
    assert (restored[50:81,290:295] == img[50:81,290:295]).all()
