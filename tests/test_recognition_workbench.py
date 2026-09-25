"""Work through the real workbench controls without an OCR service or desktop clipboard."""
import json
import time
from copy import deepcopy

import pytest

pytest.importorskip('PySide6')
from PySide6.QtCore import Qt
from PySide6.QtGui import QImage
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QMenu

from shiguang_capture.ocr.base import OCRResult
from shiguang_capture.structured import TableData
from shiguang_capture.translate import TranslationResult
from shiguang_capture.ui.result_panel import ResultPanel


@pytest.fixture
def panel(qt_session):
    widget = ResultPanel()
    image = QImage(240, 160, QImage.Format.Format_RGB32)
    image.fill(Qt.GlobalColor.white)
    widget.set_image(image)
    widget.show()
    qt_session.processEvents()
    yield widget
    widget.close()
    widget.deleteLater()


def result(text='Original text', **kwargs):
    return OCRResult(text, .95, **kwargs)


def translation(text='Original text', target='原文'):
    return TranslationResult(text, target, 'en', 'zh', 'test')


def menu_action(menu, title):
    return next(action for action in menu.actions() if action.text() == title)


def test_import_controls_and_mode_menu_dispatch_real_image(panel):
    opened, pasted, recognized = [], [], []
    panel.open_requested.connect(lambda: opened.append(True))
    panel.paste_requested.connect(lambda: pasted.append(True))
    panel._delegate = lambda image, mode: recognized.append((image, mode))
    QTest.mouseClick(panel.open_button, Qt.MouseButton.LeftButton)
    QTest.mouseClick(panel.paste_button, Qt.MouseButton.LeftButton)
    assert opened == pasted == [True]
    assert panel.open_button.isVisible() and panel.paste_button.isVisible()

    panel.show_result('ocr', result())
    menu = QMenu(panel)
    panel._populate_result_menu(menu)
    modes = menu_action(menu, '识别模式').menu()
    menu_action(modes, '代码').trigger()
    assert len(recognized) == 1
    assert recognized[0][1] == 'code'
    assert recognized[0][0].size() == panel.canvas.image.size()
    assert panel.mode_combo.currentData() == 'code'
    assert {'复制格式', '导出', '清理选中文字…', '清空会话'} <= {a.text() for a in menu.actions()}


def test_edit_during_translation_cancels_and_rejects_stale_result(panel):
    cancelled = []
    panel.cancel_requested.connect(lambda: cancelled.append(True))
    original = result()
    panel.show_result('ocr', original)
    panel._translate()
    panel.set_busy(True)
    panel.source_edit.setPlainText('Corrected text')
    assert cancelled
    assert not panel._busy
    assert panel._pending_translation is None
    panel.show_translation(original, translation())
    assert panel.source_edit.toPlainText() == 'Corrected text'
    assert not panel.target_edit.toPlainText()
    assert '原文已修改' in panel.feedback.text()


def test_typing_while_ocr_runs_keeps_edit_and_new_retry_can_complete(panel):
    cancelled = []
    panel.cancel_requested.connect(lambda: cancelled.append(True))
    panel.set_busy(True)
    panel.source_edit.setPlainText('Keep my correction')
    panel.show_result('ocr', result('Late recognition'))
    assert cancelled
    assert panel.source_edit.toPlainText() == 'Keep my correction'
    assert not panel._busy
    panel.set_busy(True)
    panel.show_result('ocr', result('New retry result'))
    assert panel.source_edit.toPlainText() == 'New retry result'


def test_clear_cancels_and_prevents_late_content_from_reappearing(panel):
    cancelled = []
    panel.cancel_requested.connect(lambda: cancelled.append(True))
    panel.show_result('ocr', result())
    panel._translate()
    panel.set_busy(True)
    panel.clear_session()
    panel.show_result('ocr', result('Late OCR output'))
    panel.show_translation(result(), translation())
    assert cancelled
    assert panel.canvas.image.isNull()
    assert not panel.source_edit.toPlainText()
    assert not panel.target_edit.toPlainText()
    assert panel.right.isHidden()
    assert not panel._busy
    assert not panel.copy_source.isEnabled()
    assert not panel.feedback.isVisible()

    image = QImage(120, 80, QImage.Format.Format_RGB32)
    image.fill(Qt.GlobalColor.white)
    panel.set_image(image)
    panel.show_result('ocr', result('New session'))
    assert panel.source_edit.toPlainText() == 'New session'


def test_copy_and_low_confidence_feedback_are_visible_then_expire(panel, monkeypatch):
    copied = []
    monkeypatch.setattr('shiguang_capture.ui.result_panel.write_text', lambda text: copied.append(text))
    panel.show_result('ocr', result(blocks=[{'text': 'Original text', 'confidence': .7,
                                           'box': [[0, 0], [100, 0], [100, 20], [0, 20]]}]))
    assert panel.feedback.isVisible()
    assert '1 处文字建议校对' in panel.feedback.text()
    assert '低置信' in panel.block_combo.itemText(1)
    panel._copy_source()
    assert copied == ['Original text']
    assert panel.feedback.isVisible()
    assert panel.feedback.text() == '已复制识别内容'
    panel.feedback_timer.start(1)
    # Shared CI runners can postpone Qt's first timer dispatch beyond 20 ms.
    deadline = time.monotonic() + 1
    while panel.feedback.isVisible() and time.monotonic() < deadline:
        QTest.qWait(10)
    assert not panel.feedback.isVisible()


def test_comparison_controls_reachable_and_table_edits_survive_translation(panel, monkeypatch):
    original = result('00123\t进行中', table=TableData([['00123', '进行中']], [[(0, 0, 30, 20), (30, 0, 50, 20)]]))
    panel.show_result('table', original)
    cell = panel.table_grid.item(0, 1)
    cell.setText('已完成')
    corrected = deepcopy(original)
    corrected.text = panel.source_edit.toPlainText()
    panel._translate()
    panel.show_translation(corrected, translation(corrected.text, '00123\tDone'))
    assert panel.table_grid.item(0, 1) is cell
    assert cell.text() == '已完成'
    assert panel.right.isVisible()
    menu = QMenu(panel)
    panel._populate_result_menu(menu)
    menu_action(menu, '左右对照').trigger()
    assert panel.comparison.orientation() == Qt.Orientation.Horizontal
    assert menu_action(menu, '复制译文').isEnabled()
    menu_action(menu, '交换原文和译文').trigger()
    assert panel.source_edit.isVisible() and panel.table_grid.isHidden()
    assert panel.source_edit.toPlainText() == '00123\tDone'
    assert panel.target_edit.toPlainText() == '00123\t已完成'
    copied = []
    monkeypatch.setattr('shiguang_capture.ui.result_panel.write_text', lambda text: copied.append(text))
    panel._copy_source()
    assert copied == ['00123\tDone']
    panel.source_edit.setPlainText('Changed after swap')
    assert not panel.target_edit.toPlainText()
    panel._restore()
    assert panel.table_grid.isVisible()
    assert panel.table_grid.item(0, 1).text() == '进行中'
    assert panel._kind == 'table'


@pytest.mark.parametrize('extension', ['txt', 'md', 'json', 'xlsx'])
def test_export_uses_edited_table_values(panel, monkeypatch, tmp_path, extension):
    panel.show_result('table', result('00123', table=TableData([['编号', '状态'], ['00123', '进行中']],
                                                            [[(0, 0, 30, 20)] * 2] * 2)))
    panel.table_grid.item(1, 1).setText('已完成')
    destination = tmp_path / f'output.{extension}'
    monkeypatch.setattr('shiguang_capture.ui.result_panel.QFileDialog.getSaveFileName',
                        lambda *args: (str(destination), ''))
    panel._export_output(extension)
    assert destination.is_file()
    if extension == 'json':
        assert json.loads(destination.read_text(encoding='utf-8'))['cells'][1] == ['00123', '已完成']
    elif extension == 'xlsx':
        import openpyxl
        sheet = openpyxl.load_workbook(destination).active
        assert sheet.cell(2, 1).value == '00123'
        assert sheet.cell(2, 2).value == '已完成'
    else:
        assert '00123' in destination.read_text(encoding='utf-8')
        assert '已完成' in destination.read_text(encoding='utf-8')
    assert panel.feedback.text() == '文件已导出'
