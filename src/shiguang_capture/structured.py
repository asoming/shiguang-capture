"""Observable layout reconstruction and safe structured output (no semantic correction)."""
from __future__ import annotations
from dataclasses import dataclass
import html
import io
import statistics
import unicodedata


@dataclass
class TableData:
    cells: list[list[str]]
    boxes: list[list[tuple[int, int, int, int]]]

    @property
    def shape(self):
        return len(self.cells), len(self.cells[0]) if self.cells else 0


def bounds(block):
    points = block.get('box', [])
    if len(points) < 4:
        return None
    xs, ys = [float(p[0]) for p in points], [float(p[1]) for p in points]
    return min(xs), min(ys), max(xs), max(ys)


def _display_width(text):
    return sum(2 if unicodedata.east_asian_width(c) in {'W', 'F'} else 1 for c in text)


def text_from_blocks(blocks):
    """Join fragments on an observed baseline, preserving all recognized glyphs."""
    positioned = [(bounds(block), block.get('text', '')) for block in blocks if bounds(block)]
    if not positioned:
        return '\n'.join(block.get('text', '') for block in blocks)
    positioned.sort(key=lambda item: ((item[0][1]+item[0][3])/2, item[0][0]))
    height = statistics.median(max(1, box[3]-box[1]) for box, text in positioned)
    rows = []
    for box, text in positioned:
        center = (box[1]+box[3])/2
        if not rows or abs(center-rows[-1][0]) > height*.45:
            rows.append((center, []))
        rows[-1][1].append((box, text))
    lines = []
    for center, entries in rows:
        line, previous = '', None
        for box, text in sorted(entries):
            if previous is not None:
                gap = box[0]-previous[2]
                # Large gaps may be separate columns, so do not imply a sentence.
                line += '\n' if gap > height*3 else (' ' if gap > height*.3 else '')
            line += text
            previous = box
        lines.append(line)
    return '\n'.join(lines)


def code_from_blocks(blocks):
    """Recover visible monospaced columns. Never repair spelling or execute output."""
    positioned = [(bounds(b), b['text']) for b in blocks if bounds(b) and b.get('text')]
    if not positioned:
        return '\n'.join(b.get('text', '') for b in blocks)
    positioned.sort(key=lambda item: (item[0][1]+item[0][3])/2)
    height = statistics.median(b[3]-b[1] for b, text in positioned)
    widths = [(b[2]-b[0])/_display_width(text) for b, text in positioned if _display_width(text) >= 4]
    char_width = statistics.median(widths) if widths else max(1, height*.55)
    left = min(b[0] for b, text in positioned)
    rows = []
    for box, text in positioned:
        cy = (box[1]+box[3])/2
        if not rows or abs(cy-rows[-1][0]) > height*.5:
            rows.append((cy, []))
        rows[-1][1].append((box, text))
    result = []
    for cy, fragments in rows:
        line = ''
        for box, text in sorted(fragments):
            column = max(0, round((box[0]-left)/char_width))
            line += ' '*max(0, column-_display_width(line)) + text
        result.append(line)
    return '\n'.join(result)


def _clusters(indices):
    groups = []
    for value in indices:
        if not groups or value > groups[-1][-1]+1:
            groups.append([int(value)])
        else:
            groups[-1].append(int(value))
    return [round(statistics.mean(group)) for group in groups]


def detect_grid(image_bytes):
    """Return full-border, unmerged grid and a line-cleaned image for OCR."""
    import cv2
    import numpy as np
    image = cv2.imdecode(np.frombuffer(image_bytes, dtype=np.uint8), cv2.IMREAD_COLOR)
    if image is None:
        raise ValueError('无法读取表格图片。')
    gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)
    mask = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY_INV+cv2.THRESH_OTSU)[1]
    h, w = gray.shape
    horizontal = cv2.morphologyEx(mask, cv2.MORPH_OPEN, np.ones((1, max(20, w//35)), np.uint8))
    vertical = cv2.morphologyEx(mask, cv2.MORPH_OPEN, np.ones((max(20, h//35), 1), np.uint8))
    grid = cv2.bitwise_or(horizontal, vertical)
    contours, _ = cv2.findContours(grid, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
    candidates = sorted((cv2.boundingRect(c) for c in contours), key=lambda r: r[2]*r[3], reverse=True)
    for x, y, width, height in candidates:
        if width < 40 or height < 30:
            continue
        region_h, region_v = horizontal[y:y+height, x:x+width], vertical[y:y+height, x:x+width]
        ys = [v+y for v in _clusters(np.flatnonzero(np.count_nonzero(region_h, axis=1) >= width*.92))]
        xs = [v+x for v in _clusters(np.flatnonzero(np.count_nonzero(region_v, axis=0) >= height*.92))]
        if len(xs) < 2 or len(ys) < 2:
            continue
        # A shorter interior divider indicates a merged/irregular cell, not an empty cell.
        min_row = min(b-a for a,b in zip(ys,ys[1:]))
        min_col = min(b-a for a,b in zip(xs,xs[1:]))
        # Count continuous dividers, not aligned glyph stems from separate rows.
        long_v = cv2.morphologyEx(region_v, cv2.MORPH_OPEN, np.ones((max(12, round(min_row*.8)),1), np.uint8))
        long_h = cv2.morphologyEx(region_h, cv2.MORPH_OPEN, np.ones((1,max(12, round(min_col*.8))), np.uint8))
        all_xs = _clusters(np.flatnonzero(np.count_nonzero(long_v, axis=0) >= min_row*.8))
        all_ys = _clusters(np.flatnonzero(np.count_nonzero(long_h, axis=1) >= min_col*.8))
        if len(all_xs) != len(xs) or len(all_ys) != len(ys):
            raise ValueError('检测到合并或不完整边框。请使用完整、无合并的简单表格，或切换普通文字识别。')
        if len(xs)-1 > 12 or len(ys)-1 > 30:
            raise ValueError('表格超过 30 行 × 12 列，请裁剪后分段处理。')
        if min(b-a for a,b in zip(xs,xs[1:])) < 12 or min(b-a for a,b in zip(ys,ys[1:])) < 12:
            continue
        boxes = [[(xs[c], ys[r], xs[c+1], ys[r+1]) for c in range(len(xs)-1)] for r in range(len(ys)-1)]
        cleaned = image.copy()
        # Erase only verified grid lines; morphology also finds stems inside letters.
        for line in xs:
            cleaned[max(0,ys[0]-2):ys[-1]+3, max(0,line-2):line+3] = 255
        for line in ys:
            cleaned[max(0,line-2):line+3, max(0,xs[0]-2):xs[-1]+3] = 255
        ok, encoded = cv2.imencode('.png', cleaned)
        if not ok:
            raise ValueError('表格预处理失败，请重试。')
        return boxes, encoded.tobytes()
    raise ValueError('没有找到完整表格边框。请使用清晰、有边框、无合并的单张表格。')


def table_from_blocks(boxes, blocks):
    cells = [['' for box in row] for row in boxes]
    fragments = [[[] for box in row] for row in boxes]
    for block in blocks:
        box = bounds(block)
        if not box:
            continue
        x, y = (box[0]+box[2])/2, (box[1]+box[3])/2
        for r, row in enumerate(boxes):
            for c, (left, top, right, bottom) in enumerate(row):
                if left <= x < right and top <= y < bottom:
                    fragments[r][c].append((box, block['text']))
    for r, row in enumerate(fragments):
        for c, entries in enumerate(row):
            cells[r][c] = text_from_blocks([
                {'text': text, 'box': [[box[0],box[1]], [box[2],box[1]], [box[2],box[3]], [box[0],box[3]]]}
                for box, text in entries])
    return TableData(cells, boxes)


def table_risks(cells):
    return [(r,c) for r,row in enumerate(cells) for c,text in enumerate(row)
            if text.lstrip().startswith(('=', '+', '-', '@')) or any(ch in text for ch in '\t\r\n')]


def markdown_table(cells):
    def escape(value):
        return html.escape(value, quote=False).replace('\\', '\\\\').replace('|', '\\|').replace('\r\n', '\n').replace('\n', '<br>')
    if not cells:
        return ''
    lines = ['| '+' | '.join(escape(v) for v in row)+' |' for row in cells]
    lines.insert(1, '| '+' | '.join('---' for v in cells[0])+' |')
    return '\n'.join(lines)


def table_clipboard(cells):
    if table_risks(cells):
        raise ValueError('表格含公式风险前缀或换行。请使用文本型 XLSX，或复制为普通文本。')
    rows = ['<tr>'+''.join('<td style="mso-number-format:\\@">'+html.escape(text)+'</td>' for text in row)+'</tr>' for row in cells]
    return '\n'.join('\t'.join(row) for row in cells), '<table>'+''.join(rows)+'</table>'


def xlsx_bytes(cells):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    import re
    for row in cells:
        for text in row:
            if len(str(text)) > 32767:
                raise ValueError('单元格超过 XLSX 的 32,767 字符限制，请缩短内容。')
            if re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', str(text)):
                raise ValueError('单元格含 XLSX 不支持的控制字符，请清理后重试。')
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '识别表格'
    for r,row in enumerate(cells,1):
        for c,text in enumerate(row,1):
            cell = sheet.cell(r,c, str(text))
            cell.data_type = 's'
            cell.number_format = '@'
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if r == 1:
                cell.font = Font(bold=True,color='FFFFFF')
                cell.fill = PatternFill('solid',fgColor='167D8D')
    from openpyxl.utils import get_column_letter
    for c in range(1,(len(cells[0]) if cells else 0)+1):
        sheet.column_dimensions[get_column_letter(c)].width = 24
    sheet.freeze_panes = 'A2'
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def code_block(text):
    longest = max((len(part) for part in __import__('re').findall(r'`+', text)), default=0)
    fence = '`'*max(3,longest+1)
    return fence+'\n'+text+'\n'+fence
